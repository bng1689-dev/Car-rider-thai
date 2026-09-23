"""สร้างไฟล์ Excel — 3 ชีต

1. เมทริกซ์กว้าง   หน้าตาคุ้นเหมือนของเดิม ใช้ดูภาพรวมเร็ว ๆ
2. ข้อกล่าวอ้าง    รูปแบบยาว พร้อมที่มา/ความเชื่อมั่น/วันที่ตรวจ ครบทุกแถว
3. งานค้าง        หมวดและรุ่นที่ยังไม่ได้ตรวจ — ระบบเดิมสร้างไม่ได้เลย

ชีต 2 คือรูปทรงเดียวกับตาราง entity-link ของ i2
(ต้นทาง · ปลายทาง · ชนิดความสัมพันธ์ · คุณสมบัติ · แหล่งที่มา · วันที่)
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from .. import view
from ..model import Reference, Resolved, Vehicle
from ..resolve import index_by_vehicle
from . import OUTPUTS

NAVY, GOLD, GRAB_GREEN, BOLT_GREEN = "16244D", "B8962E", "00803A", "1A9C5F"
WHITE, LIGHT, RED, GREY, VIOLET = "FFFFFF", "F5F7FB", "C0392B", "8A8FA8", "6B46C1"
FONT = "TH Sarabun New"

# ตัวอักษรในเซลล์บอก "ที่มา" ของข้อสรุป ไม่ใช่แค่ผ่าน/ไม่ผ่าน
CELL_TOKEN = {
    ("eligible", "official_list"): ("Y", GRAB_GREEN, True),
    ("eligible", "spec_inference"): ("I", GOLD, False),
    ("eligible", "rule_derived"): ("R", NAVY, False),
    ("eligible", "manual"): ("M", VIOLET, False),
    ("unverified", None): ("?", RED, True),
    ("unverified", "manual"): ("?", RED, True),
}
BLANK = ("–", "CCCCCC", False)

LEGEND = ("Y = ยืนยันจากรายการทางการ · I = อนุมานจาก spec · R = คำนวณจากกฎ · "
          "M = บันทึกโดยผู้จัดทำ · ? = ยังไม่ได้ตรวจ · – = ไม่ผ่าน/ไม่มีข้อมูล")


def _token(cell: Resolved | None):
    if cell is None or cell.eligibility == "ineligible":
        return BLANK
    return CELL_TOKEN.get((cell.eligibility, cell.basis), ("·", GREY, False))


def build(
    reference: Reference,
    vehicles: list[Vehicle],
    resolved: list[Resolved],
    as_of: date,
    stale_days: int,
    out_dir: Path = OUTPUTS,
) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="CCD2E0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    index = index_by_vehicle(resolved)
    categories = [c for p in reference.platforms for c in reference.categories_of(p.id)]

    wb = Workbook()

    # ───────── ชีต 1: เมทริกซ์กว้าง ─────────
    ws = wb.active
    ws.title = "เมทริกซ์"
    width = 5 + len(categories) + 2

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    title = ws.cell(row=1, column=1,
                    value="รถยนต์ไฟฟ้า (BEV) ในไทย × หมวดบริการ Grab & Bolt")
    title.font = Font(name=FONT, size=16, bold=True, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    legend = ws.cell(row=2, column=1, value=f"{LEGEND}  ·  ข้อมูล ณ {as_of.isoformat()}")
    legend.font = Font(name=FONT, size=10, italic=True, color=GOLD)
    legend.fill = PatternFill("solid", fgColor="FDF9EE")
    legend.alignment = Alignment(horizontal="center")

    headers = (["ยี่ห้อ", "รุ่น", "ตัวถัง", "ที่นั่ง", "สถานะ"]
               + [c.label for c in categories]
               + ["ตรวจล่าสุด", "หมายเหตุ"])
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=text)
        if 6 <= col < 6 + len(categories):
            category = categories[col - 6]
            fill = GRAB_GREEN if category.platform == "grab" else BOLT_GREEN
        else:
            fill = NAVY
        cell.font = Font(name=FONT, size=11, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[3].height = 34

    for offset, vehicle in enumerate(vehicles):
        row = 4 + offset
        cells = index.get(vehicle.vehicle_id, {})
        checked = sorted(
            {c.checked_at for c in cells.values() if c.checked_at}, reverse=True
        )
        values = ([vehicle.brand, vehicle.model, view.body_label(vehicle, reference),
                   vehicle.seats, view.launch_text(vehicle)]
                  + [None] * len(categories)
                  + [checked[0] if checked else "", vehicle.note or vehicle.data_issue])
        stripe = PatternFill("solid", fgColor=LIGHT) if row % 2 == 0 else None

        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            if stripe:
                cell.fill = stripe

            if 6 <= col < 6 + len(categories):
                category = categories[col - 6]
                token, color, bold = _token(cells.get(category.id))
                cell.value = token
                cell.font = Font(name=FONT, size=10.5, bold=bold, color=color)
                cell.alignment = Alignment(horizontal="center")
                detail = cells.get(category.id)
                if detail is not None and detail.eligibility != "ineligible":
                    note = f"{detail.reason}"
                    if detail.checked_at:
                        note += f"\nตรวจ {detail.checked_at}"
                    if detail.rule_conflict:
                        note += f"\n⚠ {detail.rule_conflict}"
                    cell.comment = _comment(note)
            else:
                cell.value = value
                cell.font = Font(name=FONT, size=10.5,
                                 bold=col in (1, 2),
                                 color=NAVY if col in (1, 2) else "000000")
                if col in (4, 5, 6 + len(categories)):
                    cell.alignment = Alignment(horizontal="center")

    widths = [12, 26, 16, 7, 14] + [9] * len(categories) + [12, 26]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "C4"

    # ───────── ชีต 2: ข้อกล่าวอ้างรูปแบบยาว ─────────
    ws2 = wb.create_sheet("ข้อกล่าวอ้าง")
    long_headers = ["vehicle_id", "ยี่ห้อ", "รุ่น", "แพลตฟอร์ม", "หมวด", "ชั้นย่อย",
                    "ผลการพิจารณา", "ที่มา", "ความเชื่อมั่น", "ตรวจเมื่อ",
                    "อายุ (วัน)", "เกินกำหนด", "แหล่งอ้างอิง", "เหตุผล", "กฎขัดแย้ง"]
    for col, text in enumerate(long_headers, start=1):
        cell = ws2.cell(row=1, column=col, value=text)
        cell.font = Font(name=FONT, size=11, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for offset, r in enumerate(resolved):
        ws2.append([
            r.vehicle_id, r.brand, r.model, r.platform, r.category_label, r.tier or "",
            view.ELIGIBILITY_LABEL[r.eligibility],
            view.BASIS_LABEL.get(r.basis or "", "ยังไม่ได้ตรวจ"),
            r.confidence or "", r.checked_at,
            r.staleness_days if r.staleness_days is not None else "",
            "ใช่" if r.is_stale else "", r.source_url, r.reason, r.rule_conflict,
        ])
        for col in range(1, len(long_headers) + 1):
            ws2.cell(row=2 + offset, column=col).font = Font(name=FONT, size=10)

    for col, w in enumerate([22, 12, 24, 11, 14, 9, 13, 22, 11, 12, 10, 10, 30, 40, 44], start=1):
        ws2.column_dimensions[get_column_letter(col)].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(long_headers))}{1 + len(resolved)}"

    # ───────── ชีต 3: งานค้าง ─────────
    ws3 = wb.create_sheet("งานค้าง")
    ws3.append(["หมวด", "สถานะการสำรวจ", "รุ่น", "เหตุผลที่ต้องตรวจ"])
    for col in range(1, 5):
        cell = ws3.cell(row=1, column=col)
        cell.font = Font(name=FONT, size=11, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=RED)

    # แถวที่มีข้อกล่าวอ้างแล้วแต่ยังไม่รู้ผล
    for r in (row for row in resolved if row.eligibility == "unverified"):
        category = reference.category_by_id[r.category_id]
        ws3.append([category.label, category.coverage_status,
                    f"{r.brand} {r.model}", r.reason])

    # แถวของหมวดที่ยังไม่ได้สำรวจ — ตัดรถที่มีข้อกล่าวอ้างแล้วออก
    # (มิฉะนั้นรถคันเดียวกันจะโผล่สองครั้งด้วยเหตุผลที่ขัดกันเอง)
    for gap in view.survey_gaps(reference, vehicles, resolved):
        reason = ("เข้าเกณฑ์เครื่อง แต่ยังไม่มีใครตรวจหมวดนี้กับรุ่นนี้"
                  + (f" (หมวดนี้สำรวจไปแล้ว {gap['claimed']} รุ่น)"
                     if gap["claimed"] else ""))
        for vehicle in gap["candidates"]:
            ws3.append([gap["category"].label, "not_surveyed",
                        vehicle.display_name, reason])

    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
    for col, w in enumerate([16, 16, 30, 62], start=1):
        ws3.column_dimensions[get_column_letter(col)].width = w
    ws3.freeze_panes = "A2"

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "EV-Grab-Bolt-Thailand.xlsx"
    wb.save(path)
    return path


def _comment(text: str):
    from openpyxl.comments import Comment

    comment = Comment(text, "evbuild")
    comment.width, comment.height = 320, 90
    return comment
