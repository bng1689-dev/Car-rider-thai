#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build.py — EV Thailand x Grab/Bolt category mapper
สร้างผลลัพธ์ทั้งหมดจากไฟล์ข้อมูลเดียว (ev_data.json)

Outputs -> ./outputs/
  1. EV-Grab-Bolt-Thailand.xlsx    ตารางเปรียบเทียบ (แก้ ev_data.json แล้วรันใหม่)
  2. search.html                    หน้าค้นหารุ่น -> หมวดขับ (standalone)
  3. analysis.pdf                   เอกสารวิเคราะห์ + ตาราง EV (ต้องมี Sarabun font)

วิธีใช้:
  python build.py            # สร้างทุกอย่าง
  python build.py xlsx       # เฉพาะ xlsx
  python build.py html       # เฉพาะ search html
  python build.py pdf        # เฉพาะ pdf

การอัปเดตข้อมูล:
  แก้ไข ev_data.json (เพิ่ม/ลบรุ่น หรือแก้หมวด) แล้วรัน python build.py

โครงสร้าง 1 รายการใน ev_data.json:
  brand, model, body, seats, status(sale|2026|2027), note
  ธง Grab: gcar,gprem,gsuv,gvan,gexec
  ธง Bolt: bbasic,bcomf,bxl,bprem,bgreen
  ค่า: y=ยืนยัน(Grab list) · i=inference(Bolt spec) · p=รอตรวจสอบ · -=ไม่เข้าเกณฑ์

ต้องติดตั้ง:
  pip install openpyxl weasyprint
  (PDF: ต้องมีฟอนต์ Sarabun / TH SarabunPSK ในระบบ)
"""
import json, sys, os

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "outputs")
os.makedirs(OUT, exist_ok=True)

def load():
    with open(os.path.join(HERE, "ev_data.json"), encoding="utf-8") as f:
        return json.load(f)

STMAP = {"sale": "ขายแล้ว", "2026": "เปิดตัว 2026", "2027": "คาด 2027"}
STMAP_SHORT = {"sale": "ขายแล้ว", "2026": "2026", "2027": "2027"}

# ---------- helper: readable category text ----------
def grab_cats(e):
    m = []
    if e["gcar"] == "y": m.append("Car")
    if e["gprem"] == "y": m.append("Premium")
    elif e["gprem"] == "p": m.append("Premium?")
    if e["gsuv"] == "y": m.append("SUV")
    elif e["gsuv"] == "p": m.append("SUV?")
    if e["gvan"] == "y": m.append("Van")
    if e["gexec"] == "y": m.append("Exec")
    return ", ".join(m) if m else "—"

def bolt_cats(e):
    m = []
    for k, lbl in [("bbasic","Basic"),("bcomf","Comfort"),("bxl","XL"),("bprem","Premium"),("bgreen","Green")]:
        if e[k] == "i": m.append(lbl)
    return ", ".join(m) if m else "—"

# ==================== XLSX ====================
def build_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    navy, gold, grabg, boltg, white, lightblue = "16244D","B8962E","00803A","1A9C5F","FFFFFF","F5F7FB"
    thin = Side(style="thin", color="CCD2E0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook(); ws = wb.active; ws.title = "EV x Grab-Bolt"
    ws.merge_cells("A1:P1")
    c = ws["A1"]; c.value = "รถยนต์ไฟฟ้า (BEV) ในไทย 2024–2027 × หมวดบริการ Grab & Bolt"
    c.font = Font(name="TH Sarabun New", size=16, bold=True, color=white)
    c.fill = PatternFill("solid", fgColor=navy); c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:P2")
    c = ws["A2"]; c.value = "y=ยืนยัน(Grab list) · i=inference(Bolt spec) · p=รอตรวจสอบ · - =ไม่เข้าเกณฑ์"
    c.font = Font(name="TH Sarabun New", size=10, italic=True, color=gold)
    c.fill = PatternFill("solid", fgColor="FDF9EE"); c.alignment = Alignment(horizontal="center")

    headers = ["ยี่ห้อ","รุ่น","ตัวถัง","ที่นั่ง","สถานะ","GrabCar","G-Premium","G-SUV","G-Van","G-Exec",
               "Bolt Basic","B-Comfort","B-XL","B-Premium","B-Green","หมายเหตุ"]
    hrow = 3
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=hrow, column=j, value=h)
        cell.font = Font(name="TH Sarabun New", size=11, bold=True, color=white)
        fill = grabg if 6<=j<=10 else boltg if 11<=j<=15 else navy
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[hrow].height = 34

    r = hrow + 1
    for e in data:
        vals = [e["brand"],e["model"],e["body"],e["seats"],STMAP[e["status"]],
                e["gcar"],e["gprem"],e["gsuv"],e["gvan"],e["gexec"],
                e["bbasic"],e["bcomf"],e["bxl"],e["bprem"],e["bgreen"],e["note"]]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = Font(name="TH Sarabun New", size=10.5); cell.border = border
            if j in (1,2):
                cell.font = Font(name="TH Sarabun New", size=10.5, bold=True, color=navy)
            if 6 <= j <= 15:
                cell.alignment = Alignment(horizontal="center")
                col = {"y":grabg,"i":gold,"p":"C0392B","-":"CCCCCC"}.get(v,"000000")
                bold = v == "y"
                cell.font = Font(name="TH Sarabun New", size=10.5, bold=bold, color=col)
            if j == 5:
                cell.alignment = Alignment(horizontal="center")
            if r % 2 == 0 and cell.fill.fgColor.rgb in (None, "00000000"):
                cell.fill = PatternFill("solid", fgColor=lightblue)
        r += 1

    for j, w in enumerate([12,26,16,7,13,9,10,8,8,8,10,10,7,10,9,26], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"
    path = os.path.join(OUT, "EV-Grab-Bolt-Thailand.xlsx")
    wb.save(path); print("[xlsx]", path)

# ==================== SEARCH HTML ====================
def build_html(data):
    data_js = json.dumps(data, ensure_ascii=False)
    tpl = open(os.path.join(HERE, "search_template.html"), encoding="utf-8").read()
    html = tpl.replace("/*__DATA__*/", data_js)
    path = os.path.join(OUT, "search.html")
    open(path, "w", encoding="utf-8").write(html)
    print("[html]", path)

# ==================== PDF ====================
def build_pdf(data):
    from weasyprint import HTML
    analysis = open(os.path.join(HERE, "analysis_body.html"), encoding="utf-8").read()

    # EV table rows
    rows = ""; cur = None
    stcls = {"sale":"st-sale","2026":"st-2026","2027":"st-2027"}
    for e in data:
        if e["brand"] != cur:
            cur = e["brand"]; rows += '<tr class="brandrow"><td colspan="6">%s</td></tr>' % cur
        pickup = e["gcar"] == "-"
        g = '<span class="pk">รถกระบะ — ไม่รับ</span>' if pickup else grab_cats(e)
        b = "—" if pickup else bolt_cats(e)
        rows += ('<tr><td class="mdl">%s</td><td class="c">%s</td><td class="c">%s</td>'
                 '<td class="c"><span class="stg %s">%s</span></td>'
                 '<td class="gc">%s</td><td class="bc">%s</td></tr>') % (
                 e["model"], e["body"], e["seats"], stcls[e["status"]],
                 STMAP_SHORT[e["status"]], g, b)

    ev_section = EV_STYLE + EV_SECTION_TPL % (len(data), rows)
    full = analysis.replace("<!--EV_SECTION-->", ev_section)
    path = os.path.join(OUT, "analysis.pdf")
    HTML(string=full).write_pdf(path)
    print("[pdf]", path)

EV_STYLE = """<style>
.evtable{font-size:8.5pt;} .evtable td{padding:4px 6px;}
.evtable .mdl{font-weight:700;color:#16244d;}
.evtable .gc{font-size:8pt;} .evtable .bc{font-size:8pt;color:#7a6a2e;}
.brandrow td{background:#16244d !important;color:#fff;font-weight:700;font-size:9.5pt;padding:5px 8px;}
.stg{font-size:7.5pt;padding:1px 6px;border-radius:8px;font-weight:700;}
.st-sale{background:#e0f5e9;color:#00803a;} .st-2026{background:#fdf2d9;color:#a67a00;}
.st-2027{background:#fde8d9;color:#c0560a;} .pk{color:#c0392b;font-size:8pt;}
</style>"""

EV_SECTION_TPL = """
<div class="pagebreak"></div>
<h2><span class="num">7</span>ตารางรถ EV (BEV) ในไทย 2024–2027 × หมวดบริการ</h2>
<p>รวมรถยนต์ไฟฟ้าล้วน (BEV) ที่จำหน่ายและเตรียมเปิดตัวในไทย <strong>%d รุ่น</strong>
พร้อม map หมวดที่สมัครขับได้ · เครื่องหมาย <strong>?</strong> = รุ่นใหม่ที่ยังไม่ยืนยันใน Grab list
· หมวด Bolt ทั้งหมดเป็น inference จาก spec</p>
<table class="data evtable"><thead><tr>
<th style="width:22%%">รุ่น</th><th class="c" style="width:16%%">ตัวถัง</th>
<th class="c" style="width:7%%">ที่นั่ง</th><th class="c" style="width:13%%">สถานะ</th>
<th style="width:21%%">หมวด Grab</th><th style="width:21%%">หมวด Bolt ⟡</th>
</tr></thead><tbody>%s</tbody></table>
<div class="note-box"><strong>สรุปจากตาราง:</strong> BEV ทุกรุ่น (ยกเว้นรถกระบะไฟฟ้า) ผ่าน GrabCar + Bolt Basic/Green อัตโนมัติ ·
EV ที่ยืนยัน Grab Premium: BYD Han/Seal, Tesla ทุกรุ่น, Porsche Taycan, Lexus RZ/UX, MG Maxus 9 ·
รุ่นใหม่ 2026–2027 ที่มีเครื่องหมาย "?" ต้องรอ Grab อัปเดต model list</div>
"""

if __name__ == "__main__":
    data = load()
    print(f"loaded {len(data)} EV models")
    arg = sys.argv[1] if len(sys.argv) > 1 else "all"
    if arg in ("all", "xlsx"): build_xlsx(data)
    if arg in ("all", "html"): build_html(data)
    if arg in ("all", "pdf"):
        try: build_pdf(data)
        except Exception as ex: print("[pdf] skipped:", ex)
    print("done ->", OUT)
